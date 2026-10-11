"""Shared drawing helpers for the JASIPA deck and poster (python-pptx)."""
import os
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(HERE)
SHOTS = os.path.join(DOCS, "screenshots")
FIGS = os.path.join(DOCS, "report", "images", "generated")

INK = RGBColor(0x1C, 0x20, 0x24)
MUTED = RGBColor(0x60, 0x64, 0x6C)
LINE = RGBColor(0xE4, 0xE4, 0xE7)
SOFT = RGBColor(0xF7, 0xF7, 0xF8)
ACCENT = RGBColor(0x3E, 0x63, 0xDD)
ACCENT_SOFT = RGBColor(0xED, 0xF2, 0xFE)
GREEN = RGBColor(0x18, 0x79, 0x4E)
GREEN_SOFT = RGBColor(0xE9, 0xF9, 0xEE)
AMBER = RGBColor(0xAB, 0x64, 0x00)
AMBER_SOFT = RGBColor(0xFF, 0xF7, 0xE0)
RED = RGBColor(0xCE, 0x2C, 0x31)
RED_SOFT = RGBColor(0xFF, 0xEF, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


def shot(name):
    return os.path.join(SHOTS, name)


def fig(name):
    return os.path.join(FIGS, name)


def rect(slide, x, y, w, h, fill=None, line=None, radius=False, weight=0.75):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        shp.adjustments[0] = 0.06
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(weight)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15, font=FONT, space_after=0):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, {opts}) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        parts = [(para, {})] if isinstance(para, str) else para
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=14, color=INK, gap=6, marker="•", marker_color=ACCENT):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            head, body = it
            paras.append([(marker + "  ", {"color": marker_color, "bold": True}), (head, {"bold": True}), (" " + body, {})])
        else:
            paras.append([(marker + "  ", {"color": marker_color, "bold": True}), (it, {})])
    return text(slide, x, y, w, h, paras, size=size, color=color, space_after=gap)


def picture(slide, path, x, y, w, h, border=True):
    """Fit an image inside the box, centred, preserving its aspect ratio."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    px, py = x + (w - pw) // 2, y + (h - ph) // 2
    pic = slide.shapes.add_picture(path, px, py, pw, ph)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def pill(slide, x, y, label, fill, color, size=11, w=None, h=None):
    w = w or Emu(int(Pt(size).emu * 0.62 * (len(label) + 3)))
    h = h or Emu(int(Pt(size).emu * 2.0))
    r = rect(slide, x, y, w, h, fill=fill, radius=True)
    r.adjustments[0] = 0.5
    text(slide, x, y, w, h, label, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def table(slide, x, y, w, rows, col_w, size=12, row_h=None, header_fill=SOFT, zebra=False):
    n_rows, n_cols = len(rows), len(rows[0])
    row_h = row_h or Emu(int(Pt(size).emu * 2.3))
    gt = slide.shapes.add_table(n_rows, n_cols, x, y, w, row_h * n_rows).table
    total = sum(col_w)
    for j, cw in enumerate(col_w):
        gt.columns[j].width = Emu(int(w * cw / total))
    for i, row in enumerate(rows):
        gt.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = gt.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (SOFT if zebra and i % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Emu(int(Pt(size).emu * 0.6))
            cell.margin_top = cell.margin_bottom = Emu(int(Pt(size).emu * 0.25))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            color, bold = (MUTED, True) if i == 0 else (INK, False)
            if isinstance(val, tuple):
                val, color, bold = val[0], val[1], True
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size if i else size - 1)
            r.font.bold = bold
            r.font.color.rgb = color
    return gt
