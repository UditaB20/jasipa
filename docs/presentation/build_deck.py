"""Builds docs/presentation/JASIPA_presentation.pptx from repository figures, screenshots and test results."""
import json
import os

from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from common import (ACCENT, ACCENT_SOFT, AMBER, AMBER_SOFT, DOCS, GREEN, GREEN_SOFT, HERE, INK, LINE, MUTED, RED,
                    RED_SOFT, SOFT, WHITE, bullets, fig, picture, pill, rect, shot, table, text)

SUMMARY = json.load(open(os.path.join(DOCS, "test-results", "summary.json")))
TOTALS = SUMMARY["totals"]
OVERALL = SUMMARY["overall"]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
M = Inches(0.6)
BLANK = prs.slide_layouts[6]
state = {"n": 0}


def new_slide(title, kicker=None, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    state["n"] += 1
    rect(s, 0, 0, W, H, fill=WHITE)
    rect(s, 0, 0, Inches(0.12), H, fill=ACCENT)
    if kicker:
        text(s, M, Inches(0.38), Inches(10), Inches(0.3), kicker.upper(), size=11, color=ACCENT, bold=True)
    text(s, M, Inches(0.62), Inches(12), Inches(0.7), title, size=28, bold=True)
    if subtitle:
        text(s, M, Inches(1.25), Inches(12), Inches(0.4), subtitle, size=14, color=MUTED)
    rect(s, M, H - Inches(0.5), W - 2 * M, Pt(0.75), fill=LINE)
    text(s, M, H - Inches(0.42), Inches(8), Inches(0.3), "JASIPA · Job Applicant Screening and Interview Panel Agent",
         size=10, color=MUTED)
    text(s, W - M - Inches(1), H - Inches(0.42), Inches(1), Inches(0.3), str(state["n"]), size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
    return s


def card(s, x, y, w, h, title, body, tone=ACCENT, size=13):
    rect(s, x, y, w, h, fill=SOFT, line=LINE, radius=True)
    rect(s, x, y + Inches(0.2), Inches(0.06), Inches(0.35), fill=tone)
    text(s, x + Inches(0.25), y + Inches(0.18), w - Inches(0.4), Inches(0.4), title, size=15, bold=True)
    if isinstance(body, list):
        bullets(s, x + Inches(0.25), y + Inches(0.68), w - Inches(0.45), h - Inches(0.8), body, size=size,
                gap=4 if size < 16 else 12)
    else:
        text(s, x + Inches(0.25), y + Inches(0.68), w - Inches(0.45), h - Inches(0.8), body, size=size, color=MUTED)


# 1 — Title
s = prs.slides.add_slide(BLANK)
state["n"] += 1
rect(s, 0, 0, W, H, fill=WHITE)
rect(s, 0, 0, Inches(5.2), H, fill=INK)
text(s, Inches(0.7), Inches(0.8), Inches(4), Inches(0.4), "MINOR PROJECT · 2026–27", size=12, color=ACCENT_SOFT, bold=True)
text(s, Inches(0.7), Inches(2.3), Inches(4.2), Inches(1.2), "JASIPA", size=60, color=WHITE, bold=True)
text(s, Inches(0.7), Inches(3.45), Inches(4.2), Inches(1.4),
     "Job Applicant Screening and Interview Panel Agent", size=22, color=WHITE)
text(s, Inches(0.7), Inches(5.6), Inches(4.2), Inches(1.2),
     ["[Team member names]", "Guide: [Guide name]", "Dept. of CSE (Artificial Intelligence)"], size=13,
     color=LINE, space_after=4)
text(s, Inches(5.9), Inches(1.2), Inches(6.8), Inches(0.5), "An HR agent that", size=18, color=MUTED)
for i, (h_, b_) in enumerate([
        ("Screens resumes", "against a versioned, job-anchored rubric"),
        ("Assesses skills", "with generated questions and sandboxed code"),
        ("Convenes a panel", "of three independent agents"),
        ("Checks for bias", "on every merged decision"),
        ("Leaves the decision", "to a human, with a verifiable audit trail")]):
    y = Inches(1.85) + i * Inches(0.95)
    rect(s, Inches(5.9), y, Inches(0.5), Inches(0.5), fill=ACCENT_SOFT, radius=True)
    text(s, Inches(5.9), y, Inches(0.5), Inches(0.5), str(i + 1), size=16, color=ACCENT, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(6.65), y + Inches(0.02), Inches(6), Inches(0.5),
         [[(h_, {"bold": True}), ("  " + b_, {"color": MUTED})]], size=18)

# 2 — Problem
s = new_slide("Screening is slow, inconsistent and hard to audit", "Problem",
              "Recruitment AI is classed as high-risk under the EU AI Act; NYC Local Law 144 requires bias audits.")
card(s, M, Inches(1.9), Inches(3.9), Inches(2.4), "Inconsistent evaluation",
     "Different reviewers weigh the same resume differently; criteria drift between rounds and roles.")
card(s, M + Inches(4.1), Inches(1.9), Inches(3.9), Inches(2.4), "Unchecked disparity",
     "Small applicant pools make adverse-impact checks easy to misread: no data is often reported as no bias.")
card(s, M + Inches(8.2), Inches(1.9), Inches(3.9), Inches(2.4), "Opaque automation",
     "Automated rejection without human review or an audit trail is unacceptable for hiring decisions.")
rect(s, M, Inches(4.6), W - 2 * M, Inches(1.9), fill=ACCENT_SOFT, radius=True)
text(s, M + Inches(0.35), Inches(4.8), W - 2 * M - Inches(0.7), Inches(1.6), [
    [("Problem statement", {"bold": True, "color": ACCENT})],
    "Build an HR agent that takes a candidate from resume screening through skill assessment to an interview-panel "
    "recommendation, using resume_screen, skill_assess and bias_check tools, a JD-anchored rubric, stage memory, an MCP "
    "connector to an ATS database and a swarm of evaluator agents — where no candidate is rejected without human review "
    "and every panel decision carries a logged bias check."], size=15, space_after=6)

# 3 — Objectives
s = new_slide("Objectives and status", "Objectives")
rows = [["#", "Objective", "Status"],
        ["O1", "Screen PDF, DOCX and dataset resumes against a versioned rubric", ("Done", GREEN)],
        ["O2", "Generate MCQ, coding, scenario and STAR assessments; regenerate", ("Done", GREEN)],
        ["O3", "Score code against hidden tests in an isolated sandbox", ("Done", GREEN)],
        ["O4", "Swarm of three evaluators + Panel Coordinator (LangGraph)", ("Done", GREEN)],
        ["O5", "Bias check on every panel decision, small-sample aware", ("Done · no real cohort data", AMBER)],
        ["O6", "Human decisions, two-reviewer escalation, auditable log", ("Done", GREEN)],
        ["O7", "MCP connector to the ATS database", ("Done", GREEN)],
        ["O8", "Kaggle Resume Dataset as candidate data", ("Done · 6 categories", GREEN)],
        ["O9", "Agents that learn from decisions and outcomes", ("Not implemented", RED)]]
table(s, M, Inches(1.5), W - 2 * M, rows, [0.6, 8, 3], size=14, row_h=Inches(0.52))

# 4 — Architecture
s = new_slide("System architecture", "Design", "React SPA → FastAPI → agents, sandbox and bias services → SQL database; MCP connector for agents.")
picture(s, fig("architecture.png"), M, Inches(1.75), Inches(5.6), Inches(5.1), border=False)
bullets(s, Inches(6.7), Inches(1.95), Inches(6), Inches(5), [
    ("Frontend:", "React 19 + Vite; HR workspace and candidate portal; light, minimal design system."),
    ("API:", "FastAPI with JWT, bcrypt and role-based access (HR / candidate)."),
    ("Agents:", "LangGraph state graph; each agent is a chain of strategies — LLM → rules → human expert."),
    ("Services:", "resume parser, question generator, code sandbox, fairness statistics, audit chain."),
    ("Data:", "SQLAlchemy over SQLite / PostgreSQL; versioned rubrics; candidate stage memory."),
    ("Connector:", "MCP stdio JSON-RPC + REST facade over ATS records — no decision tool exposed."),
], size=15, gap=10)

# 5 — Workflow
s = new_slide("Agentic workflow", "Orchestration", "Three evaluators run in parallel on separate evidence and never see each other's output.")
picture(s, fig("workflow.png"), M, Inches(1.75), Inches(7.6), Inches(4.9), border=False)
bullets(s, Inches(8.5), Inches(1.95), Inches(4.3), Inches(5), [
    ("Resume-Screener", "— resume vs. rubric (redacted, injection-screened)"),
    ("Skill-Assessor", "— answer keys, sandbox pass rate, open answers"),
    ("Culture-Fit", "— STAR evidence per competency"),
    ("Panel Coordinator", "— weighted merge (0.4 / 0.4 / 0.2), routes only"),
    ("Bias check", "— bound to the panel decision"),
    ("Human reviewer", "— records the decision; 2 reviewers if agents diverge ≥ 30 pts"),
], size=14, gap=9)

# 6 — Agentic components
s = new_slide("Agentic AI components", "Implementation")
rows = [["Component", "Implementation in JASIPA"],
        ["Tools", "resume_screen · skill_assess (with code sandbox) · bias_check · MCP ATS tools"],
        ["Skill", "Versioned JD-anchored rubric: resume 50/15/20/15, technical 30/40/30, competencies = 100"],
        ["Memory", "Candidate stage machine, panel history across rounds, cached question sets, audit history"],
        ["Connector", "MCP (initialize, tools/list, tools/call) over stdio, plus REST facade"],
        ["Orchestration", "LangGraph StateGraph: dispatch → 3 parallel evaluators → panel → bias"],
        ["Model", "Optional Gemini 2.5 Flash, JSON output, temperature 0.2 — deterministic fallbacks offline"],
        ["Governance", "Routing-only labels · human decision endpoint · two-reviewer rule · hash-chained log"]]
table(s, M, Inches(1.5), W - 2 * M, rows, [2.2, 10], size=14, row_h=Inches(0.6))

# 7 — Assessment + sandbox
s = new_slide("Role-specific assessments and sandboxed code", "QuestionGenerationAgent")
picture(s, shot("03-job-assessment.png"), M, Inches(1.5), Inches(6.0), Inches(4.2))
picture(s, shot("05-candidate-coding.png"), Inches(6.9), Inches(1.5), Inches(5.8), Inches(4.2))
text(s, M, Inches(5.85), Inches(6.2), Inches(1), [
    [("Generation  ", {"bold": True, "color": ACCENT}),
     ("MCQ, coding (programming roles only), scenario and STAR questions from the JD and rubric; cached per rubric version. "
      "Regenerate saves only a set that differs, and reports how many questions changed.", {})]], size=12, color=MUTED)
text(s, Inches(6.9), Inches(5.85), Inches(5.8), Inches(1), [
    [("Sandbox  ", {"bold": True, "color": ACCENT}),
     ("Fresh isolated interpreter, CPU / memory / file / process limits, audit hook blocks network, files and "
      "subprocesses. Judging happens outside the sandbox, so forged output cannot pass.", {})]], size=12, color=MUTED)

# 8 — Bias method
s = new_slide("Bias check: three honest states", "Fairness",
              "Cohort = HR-recorded monitoring label; never inferred from resumes, never shown to agents.")
for i, (lab, fill, col, body) in enumerate([
        ("Disparity detected", RED_SOFT, RED, "Impact ratio < 0.80 (four-fifths screen) with n ≥ 5 in each cohort. Fisher p shown as significant / not significant."),
        ("No indicator detected", GREEN_SOFT, GREEN, "Screen not triggered. A significant rate difference alone is a note, not a warning. Under n = 30: low-power caution."),
        ("Insufficient data", AMBER_SOFT, AMBER, "Cohort below 5, fewer than two eligible cohorts, or no selections. Never presented as fairness.")]):
    x = M + i * Inches(4.1)
    rect(s, x, Inches(1.9), Inches(3.9), Inches(2.6), fill=fill, radius=True)
    text(s, x + Inches(0.3), Inches(2.1), Inches(3.4), Inches(0.5), lab, size=18, bold=True, color=col)
    text(s, x + Inches(0.3), Inches(2.7), Inches(3.4), Inches(1.8), body, size=13, color=INK)
bullets(s, M, Inches(4.85), W - 2 * M, Inches(1.8), [
    ("Per decision:", "AI recommendation, final human decision, post-hire outcome, score-threshold gap and rationale language."),
    ("Uncertainty:", "Wilson 95% intervals; fragility flag when one outcome would flip the screen; Fisher test verified against SciPy."),
    ("Separation:", "the bias check never changes scores or routing — it informs the human reviewer."),
], size=14, gap=7)

# 9 — Oversight + screenshots
s = new_slide("Panel, bias check and human decision", "Human oversight")
picture(s, shot("06-panel-bias.png"), M, Inches(1.5), Inches(6.0), Inches(3.75))
picture(s, shot("07-decision.png"), Inches(6.75), Inches(1.5), Inches(6.0), Inches(3.75))
bullets(s, M, Inches(5.5), W - 2 * M, Inches(1.3), [
    "A decision is accepted only for the latest panel decision with a bias check, and binds both IDs to the review.",
    "Every action is appended to a SHA-256 hash chain; Audit log → Verify integrity recomputes it.",
], size=14, gap=6)

# 10 — Dataset
s = new_slide("Data: Kaggle Resume Dataset", "Dataset", "Seeded synthetic candidates were replaced by real resumes; nothing is invented.")
for i, (num, lab) in enumerate([("2,484", "resumes in Resume.csv"), ("24", "job categories"),
                                ("6", "categories with authored JDs"), ("48", "records imported by default")]):
    x = M + i * Inches(3.05)
    rect(s, x, Inches(1.9), Inches(2.85), Inches(1.4), fill=SOFT, line=LINE, radius=True)
    text(s, x + Inches(0.25), Inches(2.0), Inches(2.4), Inches(0.7), num, size=32, bold=True, color=ACCENT)
    text(s, x + Inches(0.25), Inches(2.75), Inches(2.4), Inches(0.4), lab, size=13, color=MUTED)
picture(s, shot("09-dataset-resume.png"), Inches(6.9), Inches(3.55), Inches(5.8), Inches(3.3))
bullets(s, M, Inches(3.65), Inches(6), Inches(3.3), [
    ("Categories:", "IT, Engineering, HR, Finance, Accountant, Business Development."),
    ("Structure:", "sections rebuilt from Resume_html, so the resume reads like the original."),
    ("Provenance:", "source = dataset_import, external ID, category; badge in the UI; separate from real applicants."),
    ("Honest gaps:", "no names, cohorts, answers or outcomes — so dataset runs report incomplete evidence and insufficient data."),
], size=13, gap=7)

# 11 — Testing
s = new_slide("Testing: {} / {} passed".format(OVERALL["Pass"], OVERALL["total"]), "Validation",
              "Generated from the runners' result files by scripts/run_all_tests.py — not written by hand.")
rows = [["Suite", "Total", "Pass", "Fail"]] + [[k, v["total"], (v["Pass"], GREEN), v["Fail"]] for k, v in TOTALS.items()] + \
       [[("All", INK), (OVERALL["total"], INK), (OVERALL["Pass"], GREEN), (OVERALL["Fail"], INK)]]
table(s, M, Inches(1.9), Inches(6.2), rows, [3.4, 1, 1, 1], size=15, row_h=Inches(0.6))
bullets(s, Inches(7.3), Inches(1.9), Inches(5.4), Inches(4.8), [
    ("Unit:", "agents, sandbox escapes, question generation, Fisher / Wilson values, resume parsing."),
    ("API:", "RBAC, uploads, regeneration, decisions, two reviewers, audit chain, MCP."),
    ("Component:", "IST formatting, bias labels, resume rendering, regeneration feedback."),
    ("End-to-end:", "9 Playwright journeys — job → application → sandboxed code → panel → bias → decision → audit."),
    ("Defects found:", "'C' matched inside 'C++'; regeneration reported false success; error-check ordering; sticky sidebar."),
], size=14, gap=9)

# 12 — Learning
s = new_slide("Does the agent learn? No — by design", "Honest assessment")
rows = [["Mechanism", "Present", "What it does"],
        ["Rubric / rule-based scoring", ("Yes", GREEN), "Default for every agent; deterministic and inspectable"],
        ["LLM inference", ("Optional", AMBER), "Scoring open answers, drafting questions, summaries — cannot override objective scores"],
        ["Stored history", ("Yes", GREEN), "Screenings, assessments, panels, bias checks, reviews, outcomes, audit"],
        ["Round memory", ("Context only", AMBER), "Previous panel rounds are reported to the coordinator; scores unchanged"],
        ["Retrieval", ("No", RED), "No vector store or case retrieval"],
        ["Training / feedback loop", ("No", RED), "Hiring outcomes are descriptive analytics; rubric changes are manual"]]
table(s, M, Inches(1.5), W - 2 * M, rows, [3, 1.6, 7.5], size=14, row_h=Inches(0.58))
text(s, M, Inches(5.85), W - 2 * M, Inches(0.8),
     "Why: learning from historical hiring outcomes can encode historical bias. Any feedback loop needs a validated, audited, "
     "human-approved process — stored outcomes are the foundation for it.", size=14, color=MUTED)

# 13 — Limitations / future
s = new_slide("Limitations and future scope", "Conclusion")
card(s, M, Inches(1.5), Inches(5.9), Inches(5.2), "Limitations", [
    "No learning loop from decisions or outcomes",
    "No real cohort data — bias paths validated on controlled tests",
    "Dataset has no outcomes, so accuracy vs. hiring ground truth is not measured",
    "LLM paths tested with mocks; output quality not evaluated",
    "Subprocess sandbox is defence in depth; Docker recommended",
    "Low statistical power under 30 candidates per cohort"], tone=AMBER, size=17)
card(s, M + Inches(6.2), Inches(1.5), Inches(5.9), Inches(5.2), "Future scope", [
    "Audited, human-approved rubric updates from outcomes",
    "Retrieval over past evaluations for reviewer consistency",
    "Agreement metrics against expert ratings on a labelled sample",
    "More dataset categories and sandbox languages",
    "Container / VM sandbox by default; production ATS over MCP",
    "Intersectional cohorts and power analysis"], tone=GREEN, size=17)

# 14 — Thanks
s = prs.slides.add_slide(BLANK)
state["n"] += 1
rect(s, 0, 0, W, H, fill=INK)
text(s, M, Inches(2.6), W - 2 * M, Inches(1), "Thank you", size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, M, Inches(3.8), W - 2 * M, Inches(0.6), "Questions and live demo", size=22, color=LINE, align=PP_ALIGN.CENTER)
text(s, M, Inches(5.6), W - 2 * M, Inches(0.5), "github.com/UditaB20/jasipa", size=16, color=ACCENT_SOFT,
     align=PP_ALIGN.CENTER)

out = os.path.join(HERE, "JASIPA_presentation.pptx")
prs.save(out)
print("saved", out, len(prs.slides), "slides")
