"""Builds docs/poster/JASIPA_poster.pptx (A1 portrait); export to PDF with LibreOffice."""
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "presentation"))
from pptx import Presentation  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402
from pptx.util import Inches  # noqa: E402

from common import (ACCENT, ACCENT_SOFT, AMBER, AMBER_SOFT, DOCS, GREEN, GREEN_SOFT, INK, LINE, MUTED, RED,  # noqa: E402
                    RED_SOFT, SOFT, WHITE, bullets, fig, picture, rect, shot, table, text)

HERE = os.path.dirname(os.path.abspath(__file__))
SUMMARY = json.load(open(os.path.join(DOCS, "test-results", "summary.json")))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(23.39), Inches(33.11)
W, H = prs.slide_width, prs.slide_height
s = prs.slides.add_slide(prs.slide_layouts[6])
M = Inches(0.9)
GAP = Inches(0.6)
COL = (W - 2 * M - GAP) // 2
X1, X2 = M, M + COL + GAP
BODY = 24

rect(s, 0, 0, W, H, fill=WHITE)

# Header
rect(s, 0, 0, W, Inches(4.6), fill=INK)
text(s, M, Inches(0.7), Inches(16), Inches(0.5), "MINOR PROJECT 2026–27 · DEPARTMENT OF CSE (ARTIFICIAL INTELLIGENCE)",
     size=22, color=ACCENT_SOFT, bold=True)
text(s, M, Inches(1.3), Inches(21), Inches(1.4), "JASIPA", size=110, color=WHITE, bold=True)
text(s, M, Inches(3.0), Inches(21), Inches(0.8), "Job Applicant Screening and Interview Panel Agent",
     size=48, color=WHITE)
text(s, M, Inches(3.9), Inches(21), Inches(0.5), "[Team member names] · Guide: [Guide name]", size=24, color=LINE)


def section(x, y, w, title):
    rect(s, x, y + Inches(0.08), Inches(0.12), Inches(0.55), fill=ACCENT)
    text(s, x + Inches(0.3), y, w, Inches(0.7), title, size=36, bold=True)
    return y + Inches(0.95)


# Tagline band
y = Inches(5.0)
rect(s, M, y, W - 2 * M, Inches(1.5), fill=ACCENT_SOFT, radius=True)
text(s, M + Inches(0.5), y, W - 2 * M - Inches(1), Inches(1.5),
     "Three independent agents screen, assess and evaluate each candidate; a coordinator merges their evidence; "
     "every recommendation carries a bias check and goes to a human, who records the decision in a verifiable audit trail.",
     size=28, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# Left column
y = section(X1, Inches(7.0), COL, "Problem")
bullets(s, X1, y, COL, Inches(3), [
    "Manual screening is slow and inconsistent between reviewers and rounds.",
    "Small applicant pools make adverse-impact checks easy to misread — no data is often reported as no bias.",
    "Recruitment AI is high-risk under the EU AI Act; automated rejection without human review is unacceptable.",
], size=BODY, gap=10)

y = section(X1, Inches(11.5), COL, "Architecture")
picture(s, fig("architecture.png"), X1, y, COL, Inches(8.3), border=False)

y = section(X1, Inches(21.0), COL, "Agentic workflow")
picture(s, fig("workflow.png"), X1, y, COL, Inches(4.4), border=False)
bullets(s, X1, y + Inches(4.6), COL, Inches(2.4), [
    ("Tools:", "resume_screen · skill_assess (sandboxed code) · bias_check · MCP ATS tools"),
    ("Skill:", "versioned, JD-anchored rubric with stage weights 0.4 / 0.4 / 0.2"),
    ("Escalation:", "two reviewers when agents diverge by ≥ 30 points"),
], size=20, gap=6)

# Right column
y = section(X2, Inches(7.0), COL, "Bias check: three honest states")
cw = (COL - Inches(0.5)) // 3
for i, (lab, fill, col, body) in enumerate([
        ("Disparity detected", RED_SOFT, RED, "Impact ratio < 0.80 with n ≥ 5 per cohort"),
        ("No indicator detected", GREEN_SOFT, GREEN, "Screen not triggered; significance is a note"),
        ("Insufficient data", AMBER_SOFT, AMBER, "Too few candidates — never shown as fair")]):
    x = X2 + i * (cw + Inches(0.25))
    rect(s, x, y, cw, Inches(2.4), fill=fill, radius=True)
    text(s, x + Inches(0.3), y + Inches(0.25), cw - Inches(0.6), Inches(0.9), lab, size=24, bold=True, color=col)
    text(s, x + Inches(0.3), y + Inches(1.2), cw - Inches(0.6), Inches(1.1), body, size=20)
text(s, X2, y + Inches(2.6), COL, Inches(1.0),
     "Fisher's exact test and Wilson 95% intervals quantify uncertainty. Cohorts are HR-recorded labels, never inferred "
     "and never shown to agents. The check informs the reviewer; it never changes scores or routing.", size=20, color=MUTED)

y = section(X2, Inches(11.9), COL, "System in use")
half = (COL - Inches(0.3)) // 2
picture(s, shot("03-job-assessment.png"), X2, y, half, Inches(3.4))
picture(s, shot("05-candidate-coding.png"), X2 + half + Inches(0.3), y, half, Inches(3.4))
picture(s, shot("06-panel-bias.png"), X2, y + Inches(3.6), half, Inches(3.4))
picture(s, shot("07-decision.png"), X2 + half + Inches(0.3), y + Inches(3.6), half, Inches(3.4))
text(s, X2, y + Inches(7.1), COL, Inches(0.6),
     "Generated assessment · sandboxed coding · panel with bias check · human decision", size=18, color=MUTED,
     align=PP_ALIGN.CENTER)

y = section(X2, Inches(21.0), COL, "Data and validation")
for i, (num, lab) in enumerate([("2,484", "Kaggle resumes"), ("48", "imported (6 roles)"),
                                ("{}/{}".format(SUMMARY["overall"]["Pass"], SUMMARY["overall"]["total"]), "tests passed")]):
    x = X2 + i * (cw + Inches(0.25))
    rect(s, x, y, cw, Inches(1.9), fill=SOFT, line=LINE, radius=True)
    text(s, x + Inches(0.3), y + Inches(0.15), cw, Inches(1), num, size=50, bold=True, color=ACCENT)
    text(s, x + Inches(0.3), y + Inches(1.2), cw, Inches(0.5), lab, size=20, color=MUTED)
rows = [["Suite", "Total", "Pass", "Fail"]] + [[k, v["total"], (v["Pass"], GREEN), v["Fail"]]
                                              for k, v in SUMMARY["totals"].items()]
table(s, X2, y + Inches(2.2), COL, rows, [4, 1, 1, 1], size=20, row_h=Inches(0.6))
bullets(s, X2, y + Inches(4.85), COL, Inches(2.2), [
    "Real Kaggle resumes replace synthetic seed data; dataset records are labelled and kept apart from applicants.",
    "Test tables are generated from the runners' result files — unit, API, component and 9 end-to-end journeys.",
], size=20, gap=6)

# Footer band: honest learning + limits
y = Inches(29.0)
rect(s, M, y, W - 2 * M, Inches(3.3), fill=SOFT, line=LINE, radius=True)
text(s, M + Inches(0.5), y + Inches(0.3), COL - Inches(0.5), Inches(0.6), "Does it learn?", size=30, bold=True)
text(s, M + Inches(0.5), y + Inches(1.0), COL - Inches(0.8), Inches(2.2),
     "No. Scoring is rubric-based and inspectable; an optional LLM helps with open answers, questions and summaries. "
     "Outcomes are stored for analytics only — learning from past hiring risks encoding past bias.", size=21, color=INK)
text(s, X2 + Inches(0.2), y + Inches(0.3), COL - Inches(0.5), Inches(0.6), "Limitations · next steps", size=30, bold=True)
text(s, X2 + Inches(0.2), y + Inches(1.0), COL - Inches(0.7), Inches(2.2),
     "No real cohort data or hiring outcomes; LLM output quality not evaluated; subprocess sandbox is defence in depth. "
     "Next: audited rubric updates, expert-agreement validation, container sandbox by default.", size=21, color=INK)
text(s, M, H - Inches(0.75), W - 2 * M, Inches(0.5), "github.com/UditaB20/jasipa", size=20, color=MUTED,
     align=PP_ALIGN.CENTER)

out = os.path.join(HERE, "JASIPA_poster.pptx")
prs.save(out)
print("saved", out)
